"""表格读写：Excel / CSV → 二维行；生成导入模板。

只做两件事，不碰数据库：
1. 「文件 → 二维表」（xlsx 走 openpyxl，CSV 兼容 UTF-8 与 GBK —— 旧数据大量来自
    Excel 另存的 CSV，几乎都是 GBK）；
2. 「表头 → 字段 key」的**两遍匹配**（精确 → 包含），因为老师手里的表格列名各不相同。

为什么要有别名匹配：旧应用的 `FIELD_SYNONYMS` 就是这个思路 —— 同一列可能叫
「内容 / 事项 / 具体内容」。没有它，导入就会变成「先让老师把表头改成我们想要的样子」，
而这正是「录入效率是首要瓶颈」里最该被消灭的环节。
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

from app.schemas.registry import FieldSpec, TableSpec

# 字段 key → 别人表格里可能出现的写法
ALIASES: dict[str, tuple[str, ...]] = {
    "content": ("内容", "事项", "任务", "待办", "具体内容", "正文", "话术", "模板内容"),
    "due_date": ("截止日期", "截止", "到期日", "期限", "完成日期"),
    "priority": ("优先级", "紧急程度", "重要程度"),
    "done": ("状态", "完成情况", "是否完成", "已完成"),
    "note": ("备注", "说明", "附注", "补充说明"),
    "category": ("类别", "分类", "场景", "类型"),
    "title": ("标题", "名称", "主题", "班规标题"),
    "version": ("版本", "版本号"),
    "effective_from": ("生效日期", "开始日期", "起始日期"),
    "effective_to": ("失效日期", "结束日期"),
    "use_count": ("使用次数", "次数", "引用次数"),
}

# 表头里常见的修饰，匹配前先剥掉
_NOISE = ("（必填）", "(必填)", "（选填）", "(选填)", "*", "＊", "：", ":")

XLSX_SUFFIXES = (".xlsx", ".xlsm")


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def normalize_header(value: Any) -> str:
    text = "" if value is None else str(value)
    for noise in _NOISE:
        text = text.replace(noise, "")
    return text.strip().replace(" ", "").replace("\u3000", "").lower()


def build_alias_index(spec: TableSpec) -> dict[str, str]:
    """规范化后的表头 → 字段 key。字段自己的 key 与中文名也进索引。"""
    index: dict[str, str] = {}
    for field_spec in spec.fields:
        for candidate in (field_spec.k, field_spec.label, *ALIASES.get(field_spec.k, ())):
            index.setdefault(normalize_header(candidate), field_spec.k)
    return index


def match_headers(spec: TableSpec, header_row: list[Any]) -> tuple[dict[int, str], list[str]]:
    """返回 `(列序号 → 字段 key, 未识别的表头)`。未识别列只提示、不阻断导入。"""
    index = build_alias_index(spec)
    mapping: dict[int, str] = {}
    unknown: list[str] = []
    for position, raw in enumerate(header_row):
        key = normalize_header(raw)
        if not key:
            continue
        if key in index:
            mapping[position] = index[key]
            continue
        contains = next(
            (field_key for alias, field_key in index.items() if alias and (alias in key or key in alias)),
            None,
        )
        if contains:
            mapping[position] = contains
        else:
            unknown.append(str(raw).strip())
    return mapping, unknown


def decode_text(content: bytes) -> str:
    """UTF-8 → GBK 顺序试解码，都失败才用替换字符，避免整份文件读不出来。"""
    for encoding in ("utf-8-sig", "gbk"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def read_rows(filename: str, content: bytes) -> list[list[Any]]:
    """把上传内容读成二维表（含表头行）。"""
    name = (filename or "").lower()
    return _read_xlsx(content) if name.endswith(XLSX_SUFFIXES) else _read_csv(content)


def _read_xlsx(content: bytes) -> list[list[Any]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]  # 只看第一个工作表
        # 空行**跳过**而不是截断：老师的表格中间留个空行很常见，
        # 遇到空行就 break 会把后面所有数据静默丢掉（CSV 那条路一直是跳过的，
        # 两种格式两种行为本身就是隐患）
        return [list(row) for row in sheet.iter_rows(values_only=True) if any(not is_blank(cell) for cell in row)]
    finally:
        workbook.close()


def _read_csv(content: bytes) -> list[list[str]]:
    text = decode_text(content)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel  # 嗅不出来就按标准逗号
    return [row for row in csv.reader(io.StringIO(text), dialect) if any(str(cell).strip() for cell in row)]


def _example(field_spec: FieldSpec) -> Any:
    if field_spec.type == "select":
        return field_spec.options[0] if field_spec.options else ""
    if field_spec.type == "date":
        return "2026-09-01"
    if field_spec.type == "number":
        return 0
    if field_spec.type == "checkbox":
        return "是"
    return "示例"


def build_template(spec: TableSpec) -> bytes:
    """生成导入模板：表头用中文名，第二行是示例，字段说明写进批注。"""
    editable = [field_spec for field_spec in spec.fields if field_spec.editable]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = spec.title[:31]
    sheet.append([field_spec.label for field_spec in editable])
    sheet.append([_example(field_spec) for field_spec in editable])

    for position, field_spec in enumerate(editable, start=1):
        sheet.column_dimensions[get_column_letter(position)].width = max(
            12, min(36, len(field_spec.label) * 2 + 6)
        )
        notes = [f"{field_spec.label}（{'必填' if field_spec.required else '选填'}）"]
        if field_spec.options:
            notes.append("可选值：" + "、".join(field_spec.options))
        if field_spec.type == "date":
            notes.append("格式：2026-09-01")
        if field_spec.hint:
            notes.append(field_spec.hint)
        sheet.cell(row=1, column=position).comment = Comment("\n".join(notes), "班主任工作台")

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_export(spec: TableSpec, rows: Iterable[Any]) -> bytes:
    """导出为 xlsx：列用中文字段名，日期 ISO，勾选输出「是/否」。

    输出形状与导入模板一致 —— **导出的文件改一改就能直接导回来**。
    这对「先导出、线下批量改、再导入」这个真实用法很重要。
    """
    fields = [field_spec for field_spec in spec.fields if field_spec.editable]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = spec.title[:31]
    sheet.append(["ID"] + [field_spec.label for field_spec in fields])

    for row in rows:
        values: list[Any] = [getattr(row, "id", None)]
        for field_spec in fields:
            raw = getattr(row, field_spec.k, None)
            if field_spec.type == "checkbox":
                values.append("是" if raw else "否")
            elif isinstance(raw, datetime):
                values.append(raw.strftime("%Y-%m-%d %H:%M:%S"))
            elif isinstance(raw, date):
                values.append(raw.isoformat())
            else:
                values.append(raw)
        sheet.append(values)

    sheet.column_dimensions["A"].width = 8
    for position, field_spec in enumerate(fields, start=2):
        sheet.column_dimensions[get_column_letter(position)].width = max(
            12, min(36, len(field_spec.label) * 2 + 6)
        )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
