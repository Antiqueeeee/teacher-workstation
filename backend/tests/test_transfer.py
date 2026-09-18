"""导入基础设施的端到端测试。

守的是这几条底线：
- 老师手里的表格列名五花八门 → **别名匹配**要认得出；
- 旧数据常是 Excel 另存的 GBK CSV → **解码**不能崩；
- 有脏格子时 → 预览要说清楚是哪一行哪一列，**提交要整批拒绝**（不能写进半份）；
- 重复导入同一份文件 → 不能翻倍（幂等）。
"""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

CSV_MIME = "text/csv"


def _upload(client, table: str, filename: str, content: bytes):
    return client.post(
        "/api/v1/transfer/import",
        params={"table": table},
        files={"file": (filename, content, CSV_MIME)},
    )


def _commit(client, table: str, rows: list[dict]):
    return client.post(
        "/api/v1/transfer/import/commit",
        params={"table": table},
        json={"rows": rows},
    )


def test_template_download_has_field_labels(client):
    response = client.get("/api/v1/transfer/template/todos.xlsx")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]

    sheet = load_workbook(io.BytesIO(response.content)).active
    headers = [cell.value for cell in sheet[1]]
    assert headers[:3] == ["内容", "截止日期", "优先级"]
    assert sheet["A1"].comment is not None  # 字段说明写进了批注


def test_preview_matches_aliased_headers_in_gbk_csv(client):
    csv_text = "事项,截止,紧急程度,备注\n整理班会记录,2026-10-08,高,先跟班长对一遍\n"
    response = _upload(client, "todos", "待办.csv", csv_text.encode("gbk"))
    assert response.status_code == 200, response.text
    data = response.json()["data"]

    # GBK 解码 + 别名匹配都要成立
    assert data["mapping"] == {"content": "事项", "due_date": "截止", "priority": "紧急程度", "note": "备注"}
    assert data["summary"] == {"total": 1, "ok": 1, "problem": 0, "duplicateInFile": 0}
    assert data["rows"][0]["values"]["due_date"] == "2026-10-08"


def test_unknown_header_is_reported_but_not_fatal(client):
    csv_text = "内容,这个列不认识\n顺带记一笔,x\n"
    data = _upload(client, "todos", "t.csv", csv_text.encode("utf-8")).json()["data"]
    assert data["unknownHeaders"] == ["这个列不认识"]
    assert data["summary"]["ok"] == 1  # 不因为多一列就导不进来


def test_missing_required_column_is_reported(client):
    csv_text = "备注\n随手记\n"
    data = _upload(client, "todos", "t.csv", csv_text.encode("utf-8")).json()["data"]
    assert data["missingColumns"] == ["内容"]


def test_bad_cell_is_pinpointed_with_row_and_field(client):
    csv_text = "内容,截止日期\n交材料,下周一\n"
    data = _upload(client, "todos", "t.csv", csv_text.encode("utf-8")).json()["data"]
    row = data["rows"][0]
    assert row["ok"] is False
    issue = row["issues"][0]
    assert issue["field"] == "due_date"
    assert issue["code"] == "INVALID_VALUE"
    assert "2026-09-01" in issue["message"]  # 提示里给正确格式示例


def test_duplicate_rows_in_one_file_are_flagged(client):
    csv_text = "内容,截止日期\n交材料,2026-10-01\n交材料,2026-10-01\n"
    data = _upload(client, "todos", "t.csv", csv_text.encode("utf-8")).json()["data"]
    assert data["summary"]["duplicateInFile"] == 1
    assert data["rows"][1]["issues"][0]["code"] == "DUPLICATE_IN_FILE"


def test_commit_is_all_or_nothing(client):
    rows = [
        {"content": "整批拒绝测试A", "due_date": "2026-10-02"},
        {"content": "整批拒绝测试B", "due_date": "不是日期"},
    ]
    response = _commit(client, "todos", rows)
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "ROW_VALIDATION_FAILED"

    # 关键：一行不过，另一行也不能落库
    listed = client.get("/api/v1/todos", params={"q": "整批拒绝测试"}).json()
    assert listed["meta"]["total"] == 0


def test_commit_then_reimport_is_idempotent(client):
    rows = [{"content": "幂等测试", "due_date": "2026-10-03", "priority": "中"}]

    first = _commit(client, "todos", rows).json()["data"]
    assert first["created"] == 1

    second = _commit(client, "todos", rows).json()["data"]
    assert second["created"] == 0
    assert second["skippedCount"] == 1
    assert client.get("/api/v1/todos", params={"q": "幂等测试"}).json()["meta"]["total"] == 1


def test_import_into_global_table_needs_no_class(client):
    # category / tone 的取值必须是**产品真实词表**里的（见 models/template.py 顶部说明）
    rows = [
        {
            "title": "成绩下滑沟通",
            "category": "成绩关心",
            "tone": "关切鼓励",
            "content": "您好，最近注意到……",
        }
    ]
    assert _commit(client, "templates", rows).json()["data"]["created"] == 1


def test_select_outside_real_vocabulary_is_rejected(client):
    """词表之外的取值必须被挡住 —— 否则「家长沟通」这种自造场景会悄悄混进数据。"""
    rows = [{"title": "自己编的场景", "category": "家长沟通", "content": "x"}]
    response = _commit(client, "templates", rows)
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "ROW_VALIDATION_FAILED"


# ---------------------------------------------------------------------------
# 导出：核心承诺是「导出的是**当前筛选结果**」，不是全表
# ---------------------------------------------------------------------------


def test_export_respects_current_filters(client):
    for index in (1, 2):
        client.post("/api/v1/todos", json={"content": f"导出测试-高{index}", "priority": "高"})
    client.post("/api/v1/todos", json={"content": "导出测试-低", "priority": "低"})

    response = client.get(
        "/api/v1/transfer/export/todos.xlsx",
        params={"q": "导出测试", "filter.priority": "高"},
    )
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]

    sheet = load_workbook(io.BytesIO(response.content)).active
    headers = [cell.value for cell in sheet[1]]
    assert headers[:3] == ["ID", "内容", "截止日期"]
    contents = {row[1] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert contents == {"导出测试-高1", "导出测试-高2"}  # 「低」那条没被导出去


def test_export_formats_dates_and_checkboxes_for_roundtrip(client):
    client.post(
        "/api/v1/todos",
        json={"content": "导出格式测试", "due_date": "2026-11-11", "done": True},
    )
    response = client.get("/api/v1/transfer/export/todos.xlsx", params={"q": "导出格式测试"})
    sheet = load_workbook(io.BytesIO(response.content)).active
    headers = [cell.value for cell in sheet[1]]
    row = next(iter(sheet.iter_rows(min_row=2, values_only=True)))
    data = dict(zip(headers, row))

    assert data["截止日期"] == "2026-11-11"  # ISO 字符串，不是 Excel 序列号
    # 导出用**字段名**（已完成），不是列表列名（状态）—— 字段名才是导入模板认的规范名，
    # 这样导出的文件改一改就能直接导回来
    assert data["已完成"] == "是"


def test_exported_file_can_be_imported_back(client):
    client.post("/api/v1/todos", json={"content": "回环测试", "due_date": "2026-12-01", "priority": "低"})
    exported = client.get("/api/v1/transfer/export/todos.xlsx", params={"q": "回环测试"})

    response = client.post(
        "/api/v1/transfer/import",
        params={"table": "todos"},
        files={"file": ("todos.xlsx", exported.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 200, response.text
    data = response.json()["data"]
    assert data["missingColumns"] == []
    assert data["summary"]["problem"] == 0


def test_export_rejects_unknown_table(client):
    response = client.get("/api/v1/transfer/export/nope.xlsx")
    assert response.status_code == 404
    assert response.json()["error"]["code"] == "TABLE_NOT_FOUND"


# ---------------------------------------------------------------------------
# 导入路径的回归测试（评审发现的两处「不会报错、数据却会变歪」）
# ---------------------------------------------------------------------------


def test_preview_and_commit_agree_on_defaults(client):
    """预览里补的默认值，提交后必须也是同一个 —— 曾经预览显示「中」、库里存的是空串。"""
    csv_text = "内容,截止日期\n默认值一致性测试,2026-10-09\n"
    data = _upload(client, "todos", "t.csv", csv_text.encode("utf-8")).json()["data"]
    assert data["rows"][0]["values"]["priority"] == "中"  # 预览补了默认值

    created = _commit(client, "todos", [data["rows"][0]["values"]]).json()["data"]
    assert created["created"] == 1

    listed = client.get("/api/v1/todos", params={"q": "默认值一致性测试"}).json()
    assert listed["data"][0]["priority"] == "中"  # 库里也是「中」

    # 用默认值筛选筛得到它 —— 这才说明两条路径产出的数据是同一种
    filtered = client.get(
        "/api/v1/todos", params={"q": "默认值一致性测试", "filter.priority": "中"}
    ).json()
    assert filtered["meta"]["total"] == 1


def test_xlsx_blank_row_does_not_truncate_later_rows(client):
    """表格中间的空行不能让后面的数据消失（CSV 一直是跳过，xlsx 曾经是截断）。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["内容", "截止日期"])
    sheet.append(["空行前", "2026-10-10"])
    sheet.append([None, None])
    sheet.append(["空行后", "2026-10-11"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    data = _upload(client, "todos", "带空行.xlsx", buffer.getvalue()).json()["data"]
    assert data["summary"]["total"] == 2
    assert {row["values"]["content"] for row in data["rows"]} == {"空行前", "空行后"}
