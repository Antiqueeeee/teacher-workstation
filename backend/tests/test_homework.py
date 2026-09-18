"""作业提交率的端到端测试 —— 用户明确提的痛点之一。

旧应用在这里有三套互不相通的口径：手填的 `rate`、自由文本的 `unsubmitted`（只用来出欠交榜）、
以及首页那张读不存在的字段、恒显示 0 的「作业待收」卡片。

下面的用例逐条守住「同一个数只有一个来源」这件事。
"""

from __future__ import annotations

import io

from openpyxl import load_workbook
from sqlalchemy import select

from app.models.class_ import Class
from app.models.student import Student


def _make_student(session, name: str, sno: str) -> Student:
    class_id = session.scalar(select(Class.id).where(Class.deleted_at.is_(None)).limit(1))
    student = Student(class_id=class_id, name=name, sno=sno, extra={})
    session.add(student)
    session.commit()
    return student


def _post(client, **overrides):
    payload = {
        "date": "2026-10-20",
        "subject": "数学",
        "content": "练习册 P12",
        "total": 4,
    }
    payload.update(overrides)
    return client.post("/api/v1/homework", json=payload)


def test_rate_is_computed_from_the_list(client, db_session):
    _make_student(db_session, "作业测试甲", "H9001")

    created = _post(client, unsubmitted_names="作业测试甲")
    assert created.status_code == 201, created.text
    data = created.json()["data"]

    assert data["rate"] == 75  # (4 − 1) / 4
    assert data["unsubmitted_names"] == "作业测试甲"  # 名单与提交率同源


def test_all_submitted_means_100(client):
    created = _post(client, unsubmitted_names="无", content="练习册 P13")
    assert created.json()["data"]["rate"] == 100


def test_zero_total_does_not_pretend_to_be_100(client):
    """应交 0 人时返回「—」而不是 100 —— 旧应用在这种边界上返回 100，看着像全交了。"""
    created = _post(client, total=0, content="练习册 P14")
    assert created.json()["data"]["rate"] is None


def test_editing_does_not_reset_the_total_snapshot(client, db_session):
    """编辑一条作业不该把「应交人数」重新快照成当前班级人数。

    应交人数是创建时的快照（`models/homework.py` 与变更登记表 #3 都承诺了这一点）。
    之前的钩子只要没传 total 就重算，于是转进一个学生之后，改一下「完成质量」
    就会把历史作业的应交人数改掉、提交率跟着变 —— 而且没有任何提示。
    """
    created = _post(client, unsubmitted_names="无", total=4, content="练习册 P21").json()["data"]
    assert created["total"] == 4 and created["rate"] == 100

    # 班里多一个学生（其它用例也在往同一个班加人，这里再显式加一个）
    _make_student(db_session, "快照测试甲", "H9101")

    edited = client.patch(f"/api/v1/homework/{created['id']}", json={"quality": "优"}).json()["data"]
    assert edited["total"] == 4  # 还是创建时的那个 4
    assert edited["rate"] == 100
    assert edited["quality"] == "优"

    # 显式传空值也不该把它写成 NULL（那会让 NOT NULL 列报错）或重新快照
    blanked = client.patch(f"/api/v1/homework/{created['id']}", json={"total": ""}).json()["data"]
    assert blanked["total"] == 4


def test_rate_is_recomputed_when_the_list_changes(client, db_session):
    _make_student(db_session, "作业测试乙", "H9002")
    created = _post(client, unsubmitted_names="作业测试乙", content="练习册 P15").json()["data"]
    assert created["rate"] == 75

    updated = client.patch(
        f"/api/v1/homework/{created['id']}", json={"unsubmitted_names": "作业测试乙、无此人"}
    )
    # 名单里出现认不出的人 → 明确报错，而不是把提交率算歪
    assert updated.status_code == 400
    assert "课程" not in updated.json()["error"]["message"]

    fixed = client.patch(f"/api/v1/homework/{created['id']}", json={"unsubmitted_names": "无"})
    assert fixed.json()["data"]["rate"] == 100


def test_unknown_name_is_rejected_and_nothing_is_written(client):
    response = _post(client, unsubmitted_names="查无此人甲", content="练习册 P16")
    assert response.status_code == 400
    message = response.json()["error"]["message"]
    assert "查无此人甲" in message  # 说清是谁，免得老师回去一个个对
    assert client.get("/api/v1/homework", params={"q": "练习册 P16"}).json()["meta"]["total"] == 0


def test_ambiguous_name_is_rejected(client, db_session):
    _make_student(db_session, "同名作业测试", "H9003")
    _make_student(db_session, "同名作业测试", "H9004")

    response = _post(client, unsubmitted_names="同名作业测试", content="练习册 P17")
    assert response.status_code == 400
    assert "都叫" in response.json()["error"]["message"]


def test_manual_mode_overrides_the_computed_value(client, db_session):
    _make_student(db_session, "手工模式测试", "H9005")
    created = _post(
        client,
        unsubmitted_names="手工模式测试",
        rate_mode="手工",
        rate_manual=60,
        content="练习册 P18",
    ).json()["data"]
    assert created["rate"] == 60  # 手工值生效（自动算是 75）


def test_total_defaults_to_class_size_when_omitted(client, db_session):
    _make_student(db_session, "默认人数甲", "H9104")
    _make_student(db_session, "默认人数乙", "H9105")

    created = client.post(
        "/api/v1/homework",
        json={"date": "2026-10-21", "subject": "语文", "content": "背诵篇目"},
    ).json()["data"]
    assert created["total"] == 2  # 班里就两个人，留空就按两人算（而不是 0）
    assert created["rate"] == 100  # 没填名单 = 全交


def test_manual_override_is_visible_next_to_the_computed_value(client, db_session):
    """手工覆盖时必须能看出「填的值」与「按名单算的值」不一致。

    否则同一个提交率有两个来源，而页面上完全看不出来该信哪个 ——
    这正是旧应用垮掉的方式。`rate_auto` 是给界面做对照用的派生值。
    """
    _make_student(db_session, "手工对照甲", "H9102")
    manual = _post(
        client,
        unsubmitted_names="手工对照甲",
        rate_mode="手工",
        rate_manual=60,
        content="练习册 P22",
    ).json()["data"]

    assert manual["rate"] == 60        # 生效值 = 手填
    assert manual["rate_auto"] == 75   # 按名单算 = (4 − 1) / 4

    auto = _post(client, unsubmitted_names="手工对照甲", content="练习册 P23").json()["data"]
    assert auto["rate"] == auto["rate_auto"] == 75  # 自动模式下两个值一致


def test_export_carries_both_rates(client, db_session):
    _make_student(db_session, "手工对照乙", "H9103")
    _post(
        client,
        unsubmitted_names="手工对照乙",
        rate_mode="手工",
        rate_manual=60,
        content="练习册 P24",
    )

    response = client.get("/api/v1/transfer/export/homework.xlsx", params={"q": "练习册 P24"})
    sheet = load_workbook(io.BytesIO(response.content)).active
    headers = [cell.value for cell in sheet[1]]
    row = dict(zip(headers, next(iter(sheet.iter_rows(min_row=2, values_only=True)))))
    assert row["提交率（%）"] == 60
    assert row["按名单算（%）"] == 75


def test_export_includes_rate_and_list(client, db_session):
    _make_student(db_session, "导出作业测试", "H9006")
    _post(client, unsubmitted_names="导出作业测试", content="练习册 P19")

    response = client.get("/api/v1/transfer/export/homework.xlsx", params={"q": "练习册 P19"})
    sheet = load_workbook(io.BytesIO(response.content)).active
    headers = [cell.value for cell in sheet[1]]
    row = dict(zip(headers, next(iter(sheet.iter_rows(min_row=2, values_only=True)))))
    assert row["提交率（%）"] == 75
    assert row["未交名单"] == "导出作业测试"


def test_import_creates_homework_with_computed_rate(client, db_session):
    _make_student(db_session, "导入作业测试", "H9007")
    csv_text = "布置日期,科目,应交人数,未交名单,作业内容\n2026-10-22,英语,4,导入作业测试,背诵课文\n"
    preview = client.post(
        "/api/v1/transfer/import",
        params={"table": "homework"},
        files={"file": ("h.csv", csv_text.encode("utf-8"), "text/csv")},
    ).json()["data"]
    assert preview["summary"]["problem"] == 0, preview

    rows = [row["values"] for row in preview["rows"]]
    assert client.post(
        "/api/v1/transfer/import/commit", params={"table": "homework"}, json={"rows": rows}
    ).json()["data"]["created"] == 1

    saved = client.get("/api/v1/homework", params={"q": "背诵课文"}).json()["data"][0]
    assert saved["rate"] == 75  # 导入这条路也把提交率算出来，不是留空
