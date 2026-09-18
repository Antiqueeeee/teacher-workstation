"""宿舍分布的端到端测试 —— 用户反馈「配置时有意义不明的报错」的那一块。

旧应用的问题逐条对应到这里的用例：
- 改容量用 `prompt()`、写不进去却提示成功、重开回退 8；
- 床位号「靠窗」被解析成 0，记录从网格里消失、满员误判；
- 容量调小于已住人数时界面上就少显示几个人；
- 查找不判空 → `Cannot read properties of undefined`。
"""

from __future__ import annotations

from sqlalchemy import select

from app.models.class_ import Class
from app.models.student import Student


def _class_id(session) -> int:
    return session.scalar(select(Class.id).where(Class.deleted_at.is_(None)).limit(1))


def _student(session, name: str, sno: str, *, boarding: str = "住校") -> Student:
    student = Student(class_id=_class_id(session), name=name, sno=sno, extra={"boarding": boarding})
    session.add(student)
    session.commit()
    return student


def _room(client, class_id: int, room_no: str, capacity: int = 8, building: str = "1号楼") -> dict:
    response = client.post(
        "/api/v1/dorm_rooms",
        json={"building": building, "room_no": room_no, "capacity": capacity},
        params={"classId": class_id},
    )
    assert response.status_code == 201, response.text
    return response.json()["data"]


def _bed(client, class_id: int, room_no: str, bed_no, student_name: str, **extra):
    payload = {"room_no": room_no, "bed_no": bed_no, "student_name": student_name}
    payload.update(extra)
    return client.post("/api/v1/dorm_beds", json=payload, params={"classId": class_id})


def _tree(client, class_id: int) -> dict:
    response = client.get("/api/v1/dorms/tree", params={"classId": class_id})
    assert response.status_code == 200, response.text
    return response.json()["data"]


# ---------- 容量 ----------


def test_capacity_persists_and_is_a_room_field(client, db_session):
    """容量是房间的正式字段：设完就是设完了，重开不会回退成 8（旧版会）。"""
    class_id = _class_id(db_session)
    room = _room(client, class_id, "201")
    assert room["capacity"] == 8  # 默认 8 人间

    changed = client.patch(
        f"/api/v1/dorm_rooms/{room['id']}", json={"capacity": 4}
    ).json()["data"]
    assert changed["capacity"] == 4

    again = client.get(f"/api/v1/dorm_rooms/{room['id']}").json()["data"]
    assert again["capacity"] == 4  # 不是 8


def test_capacity_cannot_drop_below_occupancy(client, db_session):
    """容量调小于已住人数 → 拒绝，并列出需要先搬走谁（旧版界面上直接少显示人）。"""
    class_id = _class_id(db_session)
    room = _room(client, class_id, "202", capacity=4)
    for index, name in enumerate(("容量测试甲", "容量测试乙", "容量测试丙"), start=1):
        _student(db_session, name, f"D90{index:02d}")
        assert _bed(client, class_id, "202", index, name).status_code == 201

    response = client.patch(f"/api/v1/dorm_rooms/{room['id']}", json={"capacity": 2})
    assert response.status_code == 400, response.text
    error = response.json()["error"]
    assert error["code"] == "DORM_CAPACITY_BELOW_OCCUPIED"
    assert "住着 3 个人" in error["message"]
    assert "容量测试丙" in error["message"]  # 说清要搬谁

    # 拒绝之后容量没变
    assert client.get(f"/api/v1/dorm_rooms/{room['id']}").json()["data"]["capacity"] == 4


def test_capacity_cannot_hide_beds_that_are_beyond_it(client, db_session):
    """容量调小必须按**最大床位号**判，不是按人数。

    8 人间里只住了 5、6 号床两个人 → 容量改成 3 时，如果只看人数（2 ≤ 3）就会放过，
    而看板只画 1..capacity 这些格子 —— 那两个人从此既不在看板上、也不在未分配名单里
    （未分配按「有没有 student_id」判定，他们已经被判成「已就座」），看起来就是「人少了」。
    """
    class_id = _class_id(db_session)
    room = _room(client, class_id, "207", capacity=8)
    for bed_no, name in ((5, "高号床位甲"), (6, "高号床位乙")):
        _student(db_session, name, f"D92{bed_no:02d}")
        assert _bed(client, class_id, "207", bed_no, name).status_code == 201

    response = client.patch(f"/api/v1/dorm_rooms/{room['id']}", json={"capacity": 3})
    assert response.status_code == 400, response.text
    error = response.json()["error"]
    assert error["code"] == "DORM_CAPACITY_BELOW_OCCUPIED"
    assert "超过" in error["message"] and "高号床位甲" in error["message"]
    assert error["detail"]["beyond"] == [5, 6]

    # 调到 6（住得下最大床位号）就能过
    assert client.patch(f"/api/v1/dorm_rooms/{room['id']}", json={"capacity": 6}).status_code == 200
    beds = [bed for bed in _tree(client, class_id)["rooms"] if bed["roomNo"] == "207"][0]["beds"]
    assert [bed["studentName"] for bed in beds if bed["studentName"]] == ["高号床位甲", "高号床位乙"]


def test_capacity_out_of_range_is_rejected(client, db_session):
    class_id = _class_id(db_session)
    room = _room(client, class_id, "203")
    response = client.patch(f"/api/v1/dorm_rooms/{room['id']}", json={"capacity": 0})
    assert response.status_code == 400
    assert "容量要在 1–" in response.json()["error"]["message"]


def test_duplicate_room_is_refused(client, db_session):
    """同一楼栋同房号不能建两次（旧应用允许，于是同一个房间被拆成两张卡）。"""
    class_id = _class_id(db_session)
    _room(client, class_id, "204")
    again = client.post(
        "/api/v1/dorm_rooms",
        json={"building": "1号楼", "room_no": "204", "capacity": 6},
        params={"classId": class_id},
    )
    assert again.status_code == 400
    assert "已经在这份名单里" in again.json()["error"]["message"]


# ---------- 床位分配 ----------


def test_bed_taken_reports_who_is_in_it(client, db_session):
    """床位已占用：告诉老师住的是谁（旧应用是「满员误判」和记录消失）。"""
    class_id = _class_id(db_session)
    _room(client, class_id, "301", capacity=2)
    _student(db_session, "占床测试甲", "D9101")
    _student(db_session, "占床测试乙", "D9102")
    assert _bed(client, class_id, "301", 1, "占床测试甲").status_code == 201

    response = _bed(client, class_id, "301", "1号床", "占床测试乙")
    assert response.status_code == 400, response.text
    error = response.json()["error"]
    assert error["code"] == "DORM_BED_TAKEN"
    assert "占床测试甲" in error["message"]


def test_one_student_one_bed(client, db_session):
    """一个学生只能有一张床：重复分配时告诉他现在住哪。"""
    class_id = _class_id(db_session)
    _room(client, class_id, "302", capacity=4)
    _student(db_session, "一床测试甲", "D9103")
    assert _bed(client, class_id, "302", 1, "一床测试甲").status_code == 201

    response = _bed(client, class_id, "302", 2, "一床测试甲")
    assert response.status_code == 400, response.text
    error = response.json()["error"]
    assert error["code"] == "DORM_STUDENT_ALREADY_ASSIGNED"
    assert "1 号床" in error["message"]


def test_bed_number_is_normalized_and_out_of_range_is_refused(client, db_session):
    """「1号床」「01」「1」是同一张床；超出容量的床位号要说清这间是几人间。"""
    class_id = _class_id(db_session)
    _room(client, class_id, "303", capacity=2)
    _student(db_session, "床号测试甲", "D9104")
    _student(db_session, "床号测试乙", "D9105")

    created = _bed(client, class_id, "303", "01", "床号测试甲")
    assert created.status_code == 201, created.text
    assert created.json()["data"]["bed_no"] == 1  # 「01」归一化成 1

    taken = _bed(client, class_id, "303", "1号床", "床号测试乙")
    assert taken.status_code == 400
    assert taken.json()["error"]["code"] == "DORM_BED_TAKEN"

    beyond = _bed(client, class_id, "303", 5, "床号测试乙")
    assert beyond.status_code == 400
    error = beyond.json()["error"]
    assert error["code"] == "DORM_BED_OUT_OF_RANGE"
    assert "2 人间" in error["message"] and "没有 5 号床" in error["message"]


def test_unparseable_bed_number_is_refused(client, db_session):
    """「靠窗」这类写法要报错 —— 旧应用把它解析成 0，那条记录从此在网格里看不见。"""
    class_id = _class_id(db_session)
    _room(client, class_id, "304")
    _student(db_session, "床号异常甲", "D9106")

    response = _bed(client, class_id, "304", "靠窗", "床号异常甲")
    assert response.status_code == 400, response.text
    assert "认不出" in response.json()["error"]["message"]
    assert "1 或 1号床" in response.json()["error"]["message"]


def test_room_is_created_on_demand_by_import(client, db_session):
    """导入时房间不存在就自动建出来 —— 老师的表里只有「楼栋/房号/床号/姓名」。"""
    class_id = _class_id(db_session)
    _student(db_session, "自动建房甲", "D9107")
    csv_text = "楼栋,房号,床位号,学生\n2号楼,105,2号床,自动建房甲\n"

    preview = client.post(
        "/api/v1/transfer/import",
        params={"table": "dorm_beds"},
        files={"file": ("d.csv", csv_text.encode("utf-8"), "text/csv")},
    ).json()["data"]
    assert preview["summary"]["problem"] == 0, preview

    values = [row["values"] for row in preview["rows"]]
    assert client.post(
        "/api/v1/transfer/import/commit",
        params={"table": "dorm_beds", "classId": class_id},
        json={"rows": values},
    ).json()["data"]["created"] == 1

    tree = _tree(client, class_id)
    room = [item for item in tree["rooms"] if item["roomNo"] == "105"][0]
    assert room["building"] == "2号楼"
    assert room["capacity"] == 8  # 自动建出来的房间用默认容量
    assert room["beds"][1]["studentName"] == "自动建房甲"  # 2 号床


def test_import_preview_reports_unparseable_bed(client, db_session):
    """「靠窗」在**导入预览**里就进冲突报告，不用等到提交才失败。"""
    class_id = _class_id(db_session)
    _student(db_session, "导入床号甲", "D9108")
    csv_text = "楼栋,房号,床位号,学生\n3号楼,201,靠窗,导入床号甲\n"

    preview = client.post(
        "/api/v1/transfer/import",
        params={"table": "dorm_beds"},
        files={"file": ("d.csv", csv_text.encode("utf-8"), "text/csv")},
    ).json()["data"]

    assert preview["summary"]["problem"] == 1, preview
    row = preview["rows"][0]
    assert row["ok"] is False
    assert "认不出" in row["issues"][0]["message"]
    assert row["issues"][0]["field"] == "bed_no"


def test_import_two_rows_same_bed_is_caught_in_the_preview(client, db_session):
    """两行指到同一个床位（「1号床」与「01」）→ **预览阶段**就认成重复并跳过第 2 行。

    判重键是（房间, 床位号），而床位号在预览解析时已经归一化成整数，所以这两种写法
    是同一个床位。原先按学生姓名判重，两行都显示正常、提交时才整批回滚 —— 老师白填一次。
    """
    class_id = _class_id(db_session)
    _student(db_session, "同床导入甲", "D9109")
    _student(db_session, "同床导入乙", "D9110")
    csv_text = "楼栋,房号,床位号,学生\n4号楼,101,1号床,同床导入甲\n4号楼,101,01,同床导入乙\n"

    preview = client.post(
        "/api/v1/transfer/import",
        params={"table": "dorm_beds"},
        files={"file": ("d.csv", csv_text.encode("utf-8"), "text/csv")},
    ).json()["data"]
    assert preview["summary"]["problem"] == 1, preview
    assert "重复" in preview["rows"][1]["issues"][0]["message"]

    values = [row["values"] for row in preview["rows"]]
    committed = client.post(
        "/api/v1/transfer/import/commit",
        params={"table": "dorm_beds", "classId": class_id},
        json={"rows": values},
    ).json()["data"]
    assert committed["created"] == 1
    assert committed["skippedCount"] == 1

    rows = client.get("/api/v1/dorm_beds", params={"classId": class_id}).json()["data"]
    assert len(rows) == 1
    assert rows[0]["student_name"] == "同床导入甲"
    assert rows[0]["bed_no"] == 1


# ---------- 编辑与腾空 ----------


def test_editing_leader_does_not_clear_the_student(client, db_session):
    """只改「寝室长」不该把学生从床上清掉（边界：钩子不能把没传的字段当作空）。"""
    class_id = _class_id(db_session)
    _room(client, class_id, "305")
    _student(db_session, "寝室长测试", "D9111")
    bed = _bed(client, class_id, "305", 3, "寝室长测试").json()["data"]

    updated = client.patch(f"/api/v1/dorm_beds/{bed['id']}", json={"leader": True}).json()["data"]
    assert updated["leader"] is True
    assert updated["student_name"] == "寝室长测试"
    assert updated["student_id"]


def test_clearing_the_name_frees_the_bed(client, db_session):
    class_id = _class_id(db_session)
    _room(client, class_id, "306", capacity=2)
    _student(db_session, "腾床测试甲", "D9112")
    bed = _bed(client, class_id, "306", 1, "腾床测试甲").json()["data"]

    updated = client.patch(
        f"/api/v1/dorm_beds/{bed['id']}", json={"student_name": ""}
    ).json()["data"]
    assert updated["student_id"] is None

    room = [item for item in _tree(client, class_id)["rooms"] if item["roomNo"] == "306"][0]
    assert room["occupied"] == 0
    assert room["beds"][0]["studentName"] == ""


def test_deleting_a_bed_frees_the_slot(client, db_session):
    """床位腾出来就是删掉这一行 —— 而且没有软删除（否则同一床位再也分配不进去）。"""
    class_id = _class_id(db_session)
    _room(client, class_id, "307", capacity=2)
    _student(db_session, "删除床位甲", "D9113")
    _student(db_session, "删除床位乙", "D9114")
    bed = _bed(client, class_id, "307", 1, "删除床位甲").json()["data"]

    assert client.delete(f"/api/v1/dorm_beds/{bed['id']}").status_code == 200
    after = client.get("/api/v1/dorm_beds", params={"classId": class_id}).json()
    assert after["meta"]["total"] == 0
    assert after["data"] == []

    # 同一个床位可以再分配给别人
    again = _bed(client, class_id, "307", 1, "删除床位乙")
    assert again.status_code == 201, again.text


# ---------- 看板视图 ----------


def test_tree_shows_slots_occupancy_and_unassigned(client, db_session):
    class_id = _class_id(db_session)
    _room(client, class_id, "401", capacity=3)
    _student(db_session, "看板测试甲", "D9115")
    _student(db_session, "看板测试乙", "D9116", boarding="走读")
    assert _bed(client, class_id, "401", 2, "看板测试甲").status_code == 201

    tree = _tree(client, class_id)
    room = [item for item in tree["rooms"] if item["roomNo"] == "401"][0]
    assert room["capacity"] == 3 and room["occupied"] == 1 and room["full"] is False
    assert [bed["bedNo"] for bed in room["beds"]] == [1, 2, 3]  # 空位也在，界面直接渲染
    assert room["beds"][1]["studentName"] == "看板测试甲"
    assert room["beds"][0]["bedId"] is None  # 空床位没有记录

    # 未分配名单只收「住校但没床位」的人
    assert [item["studentName"] for item in tree["unassigned"]] == []

    _student(db_session, "看板测试丙", "D9117")
    assert [item["studentName"] for item in _tree(client, class_id)["unassigned"]] == ["看板测试丙"]


def test_student_deleted_from_roster_flags_an_orphan_bed(client, db_session):
    """学生从档案里删掉后，床位还占着 —— 如实标出来，让老师点一下腾空。"""
    from app.db.base import utcnow

    class_id = _class_id(db_session)
    _room(client, class_id, "402")
    student = _student(db_session, "孤儿床位甲", "D9118")
    _bed(client, class_id, "402", 1, "孤儿床位甲")

    student.deleted_at = utcnow()
    db_session.commit()

    room = [item for item in _tree(client, class_id)["rooms"] if item["roomNo"] == "402"][0]
    assert room["beds"][0]["orphan"] is True
    assert room["beds"][0]["studentName"] == "孤儿床位甲"


def test_bed_without_building_uses_the_only_matching_room(client, db_session):
    """床位表只写房号（老师的表常常这样）时，用那个房号唯一对应的房间，不凭空多建一间。"""
    class_id = _class_id(db_session)
    _room(client, class_id, "501", capacity=4)
    _student(db_session, "无楼栋甲", "D9119")

    assert _bed(client, class_id, "501", 1, "无楼栋甲").status_code == 201

    rooms = [item for item in _tree(client, class_id)["rooms"] if item["roomNo"] == "501"]
    assert len(rooms) == 1  # 不是两间（「1号楼 501」和「无楼栋 501」）
    assert rooms[0]["building"] == "1号楼"
    assert rooms[0]["occupied"] == 1


def test_ambiguous_room_number_asks_for_the_building(client, db_session):
    """房号在两个楼栋都有时明确要求填楼栋 —— 猜楼的后果比报错糟糕。"""
    class_id = _class_id(db_session)
    _room(client, class_id, "601", building="1号楼")
    _room(client, class_id, "601", building="2号楼")
    _student(db_session, "歧义房号甲", "D9120")

    response = _bed(client, class_id, "601", 1, "歧义房号甲")
    assert response.status_code == 400, response.text
    assert "请填上楼栋" in response.json()["error"]["message"]


def test_export_carries_the_flat_room_and_bed_columns(client, db_session):
    """导出仍是「楼栋/房号/床位号/学生」的扁平表 —— 老师的旧表就是这个形状，
    导出去改一改要能直接导回来。"""
    import io

    from openpyxl import load_workbook

    class_id = _class_id(db_session)
    _room(client, class_id, "502", capacity=4)
    _student(db_session, "导出床位甲", "D9121")
    assert _bed(client, class_id, "502", 2, "导出床位甲").status_code == 201

    response = client.get(
        "/api/v1/transfer/export/dorm_beds.xlsx",
        params={"classId": class_id, "q": "导出床位甲"},
    )
    sheet = load_workbook(io.BytesIO(response.content)).active
    headers = [cell.value for cell in sheet[1]]
    row = dict(zip(headers, next(iter(sheet.iter_rows(min_row=2, values_only=True)))))
    assert row["楼栋"] == "1号楼"
    assert row["房号"] == "502"
    assert row["床位号"] == 2
    assert row["学生"] == "导出床位甲"


def test_missing_room_or_bed_reports_not_found(client, db_session):
    """按 id 找不到就明说「不存在」，不再有 `Cannot read properties of undefined`。"""
    response = client.get("/api/v1/dorm_rooms/999999")
    assert response.status_code == 404
    assert "不存在" in response.json()["error"]["message"]

    response = client.get("/api/v1/dorm_beds/999999")
    assert response.status_code == 404
    assert "不存在" in response.json()["error"]["message"]


def test_unknown_student_is_reported_not_guessed(client, db_session):
    class_id = _class_id(db_session)
    _room(client, class_id, "403")
    response = _bed(client, class_id, "403", 1, "宿舍查无此人")
    assert response.status_code == 400
    assert "宿舍查无此人" in response.json()["error"]["message"]


# ---------- 软删除与唯一约束（评审实测的两条 500） ----------


def test_room_can_be_recreated_after_deletion(client, db_session):
    """删掉一间房之后再建同楼栋同房号的房 —— 必须能建。

    房间是软删除的，而 (class_id, building, room_no) 原来是普通唯一约束，
    幽灵行占着键：再建一间直接撞约束报 500（出口被自己堵死）。
    现在它是**部分**唯一索引，只约束未删除的房间。
    """
    class_id = _class_id(db_session)
    room = _room(client, class_id, "900")
    assert client.delete(f"/api/v1/dorm_rooms/{room['id']}").status_code == 200

    again = client.post(
        "/api/v1/dorm_rooms",
        json={"building": "1号楼", "room_no": "900", "capacity": 4},
        params={"classId": class_id},
    )
    assert again.status_code == 201, again.text
    assert again.json()["data"]["capacity"] == 4


def test_renaming_a_room_into_a_taken_number_is_refused(client, db_session):
    """把 A2 改成已存在的 A1 → 一句说得清的话，而不是写库时才炸的唯一约束。"""
    class_id = _class_id(db_session)
    _room(client, class_id, "A1")
    second = _room(client, class_id, "A2")

    response = client.patch(f"/api/v1/dorm_rooms/{second['id']}", json={"room_no": "A1"})
    assert response.status_code == 400, response.text
    assert "已经在这份名单里" in response.json()["error"]["message"]

    # 改成没被占的房号则正常
    assert client.patch(f"/api/v1/dorm_rooms/{second['id']}", json={"room_no": "A3"}).status_code == 200


def test_room_with_occupants_cannot_be_deleted(client, db_session):
    """里面还住着人时不能删房间 —— 删了那间房的人会从所有视图里消失。

    房间软删除、床位不跟着删，看板按房间画、未分配名单按「有没有 student_id」判定，
    两边都看不到他；想给他排到别处还会撞同名房间的唯一索引报 500（评审实测）。
    """
    class_id = _class_id(db_session)
    room = _room(client, class_id, "C1", capacity=4)
    _student(db_session, "删房测试甲", "D9301")
    assert _bed(client, class_id, "C1", 1, "删房测试甲").status_code == 201

    response = client.delete(f"/api/v1/dorm_rooms/{room['id']}")
    assert response.status_code == 400, response.text
    assert "还住着 1 个人" in response.json()["error"]["message"]
    assert "删房测试甲" in response.json()["error"]["message"]

    # 房间里的人还在看板上
    board = [item for item in _tree(client, class_id)["rooms"] if item["roomNo"] == "C1"][0]
    assert board["beds"][0]["studentName"] == "删房测试甲"

    # 腾空之后就能删了
    bed_id = board["beds"][0]["bedId"]
    assert client.delete(f"/api/v1/dorm_beds/{bed_id}").status_code == 200
    assert client.delete(f"/api/v1/dorm_rooms/{room['id']}").status_code == 200
