"""出勤口径与点名的端到端测试。

旧应用在这里的毛病是「三处出勤率、口径各不同、边界假装 100%」，
以及点名白名单漏了早退导致的记录残留（`04` 缺陷清单 #8、#13）。
下面的用例逐条守住新口径，并复现改造前会失败的场景。
"""

from __future__ import annotations

import io

from openpyxl import load_workbook
from sqlalchemy import select

from app.db.base import utcnow
from app.models.class_ import Class
from app.models.student import Student
from app.services.roster import list_class_students


def _class_id(session) -> int:
    return session.scalar(select(Class.id).where(Class.deleted_at.is_(None)).limit(1))


def _student(session, name: str, sno: str) -> Student:
    student = Student(class_id=_class_id(session), name=name, sno=sno, extra={})
    session.add(student)
    session.commit()
    return student


def _roster_size(session, class_id: int) -> int:
    """全班人数 —— 出勤率的分母，直接取与后端同一个来源，免得测试自己算一遍。"""
    return len(list_class_students(session, class_id))


def _post_attendance(client, class_id: int, **overrides):
    payload = {"date": "2026-11-02", "student_name": "考勤测试甲", "type": "病假"}
    payload.update(overrides)
    return client.post("/api/v1/attendance", json=payload, params={"classId": class_id})


def _day(client, class_id: int, day: str) -> dict:
    response = client.get("/api/v1/attendance/day", params={"date": day, "classId": class_id})
    assert response.status_code == 200, response.text
    return response.json()["data"]


def _roll_call(client, class_id: int, day: str, entries: list[dict]):
    return client.put(
        "/api/v1/attendance/day",
        json={"date": day, "classId": class_id, "entries": entries},
    )


def _summary(client, class_id: int, start: str, end: str | None = None) -> dict:
    response = client.get(
        "/api/v1/attendance/summary",
        params={"from": start, "to": end or start, "classId": class_id},
    )
    assert response.status_code == 200, response.text
    return response.json()["data"]


# ---------- 出勤率口径 ----------


def test_rate_counts_students_not_records(client, db_session):
    """出勤率的分母与分子都按**人**算。

    「一个学生一天只能有一条记录」由唯一约束保证（见下一条用例），
    这条守的是分子：按去重学生数算，而不是按记录条数。
    """
    class_id = _class_id(db_session)
    a = _student(db_session, "考勤口径甲", "A9001")
    b = _student(db_session, "考勤口径乙", "A9002")

    for student, kind in ((a, "病假"), (b, "事假")):
        response = _post_attendance(
            client, class_id, date="2026-11-02", student_name=student.name, type=kind
        )
        assert response.status_code == 201, response.text

    data = _summary(client, class_id, "2026-11-02")
    expected = _roster_size(db_session, class_id)
    assert data["expected"] == expected
    assert data["absent"] == 2  # 按人算，不是按条算
    assert data["rate"] == round((expected - 2) / expected * 100)


def test_same_name_students_are_counted_one_by_one(client, db_session):
    """同名学生在出勤率里必须算**两个人** —— 去重按 student_id，不按姓名。

    按姓名去重时，班上有两个「张伟」会让出勤率偏高、缺席名单少一个人。
    这条走**点名**路径（按 studentId 提交）：出勤表单按姓名登记，同名本来就会被
    钩子拦下来报「都叫」，所以同名真正能同时登记进来的入口是点名。
    """
    class_id = _class_id(db_session)
    twin_a = _student(db_session, "考勤同名丙", "A9021")
    twin_b = _student(db_session, "考勤同名丙", "A9022")
    other = _student(db_session, "考勤同名丁", "A9023")
    day = "2026-11-20"

    response = _roll_call(
        client,
        class_id,
        day,
        [
            {"studentId": twin_a.id, "type": "旷课"},
            {"studentId": twin_b.id, "type": "旷课"},
            {"studentId": other.id, "type": "旷课"},
        ],
    )
    assert response.status_code == 200, response.text

    data = _summary(client, class_id, day)
    expected = _roster_size(db_session, class_id)
    assert data["absent"] == 3  # 不是 2
    assert data["days"][0]["absentStudents"].count("考勤同名丙") == 2
    assert data["rate"] == round((expected - 3) / expected * 100)


def test_second_record_for_the_same_student_and_day_is_rejected(client, db_session):
    """一天只能有一条：结构上有唯一约束，但老师该看到一句说得清的中文。"""
    class_id = _class_id(db_session)
    student = _student(db_session, "考勤重复甲", "A9003")
    day = "2026-11-03"

    assert _post_attendance(
        client, class_id, date=day, student_name=student.name, type="病假"
    ).status_code == 201
    again = _post_attendance(client, class_id, date=day, student_name=student.name, type="事假")

    assert again.status_code == 400, again.text
    assert "已经有一条" in again.json()["error"]["message"]


def test_late_and_early_leave_do_not_lower_the_rate(client, db_session):
    """迟到/早退单列指标，不进缺席率 —— 与旧应用的三类型过滤一致，但现在是常量而非硬编码。"""
    class_id = _class_id(db_session)
    a = _student(db_session, "考勤迟到甲", "A9004")
    b = _student(db_session, "考勤迟到乙", "A9005")
    day = "2026-11-04"

    _post_attendance(client, class_id, date=day, student_name=a.name, type="迟到")
    _post_attendance(client, class_id, date=day, student_name=b.name, type="早退")

    data = _summary(client, class_id, day)
    assert data["absent"] == 0
    assert data["late"] == 1
    assert data["early"] == 1
    assert data["rate"] == 100  # 迟到/早退不拉低出勤率


def test_unregistered_day_is_not_one_hundred_percent(client, db_session):
    """没登记考勤的那天返回「—」，不是 100% —— 旧应用两者长得一模一样。"""
    class_id = _class_id(db_session)
    data = _summary(client, class_id, "2026-11-06")

    assert data["rate"] is None
    assert data["registeredDays"] == 0
    assert data["days"][0]["registered"] is False
    assert data["days"][0]["rate"] is None  # 逐天数据里也不能冒出 100%（趋势图读它）
    assert data["days"][0]["expected"] == _roster_size(db_session, class_id)


def test_range_rate_only_uses_registered_days(client, db_session):
    """未登记的日子不进分母：把它当全员出勤只会得出一个虚高的数。"""
    class_id = _class_id(db_session)
    student = _student(db_session, "考勤区间甲", "A9006")
    _post_attendance(client, class_id, date="2026-11-07", student_name=student.name, type="病假")

    data = _summary(client, class_id, "2026-11-07", "2026-11-09")
    expected = _roster_size(db_session, class_id)
    assert data["totalDays"] == 3
    assert data["registeredDays"] == 1
    assert data["expected"] == expected          # 不是 3 × 全班人数
    assert data["absent"] == 1
    assert data["rate"] == round((expected - 1) / expected * 100)


# ---------- 点名：整体覆盖 ----------


def test_repeated_roll_call_replaces_the_whole_day(client, db_session):
    """重复提交点名不残留 —— 旧应用的点名白名单漏了「早退」，改回正常也删不掉。"""
    class_id = _class_id(db_session)
    student = _student(db_session, "考勤点名甲", "A9007")
    day = "2026-11-10"

    first = _roll_call(
        client, class_id, day, [{"studentId": student.id, "type": "早退"}]
    )
    assert first.status_code == 200, first.text
    assert first.json()["data"]["summary"]["early"] == 1
    assert first.json()["data"]["summary"]["rate"] == 100

    # 第二次提交里这名学生回到「正常」
    second = _roll_call(client, class_id, day, [])
    assert second.status_code == 200, second.text
    summary = second.json()["data"]["summary"]
    assert summary["early"] == 0
    assert summary["registered"] is False  # 一条不剩
    assert second.json()["data"]["changes"]["removed"] == 1

    assert client.get("/api/v1/attendance", params={"classId": class_id, "q": "考勤点名甲"}).json()[
        "meta"
    ]["total"] == 0


def test_roll_call_keeps_handwritten_reason_and_note(client, db_session):
    """状态没变的那条一个字都不动 —— 旧应用「删掉重建」会把老师填的处理情况抹掉。"""
    class_id = _class_id(db_session)
    student = _student(db_session, "考勤留档甲", "A9008")
    day = "2026-11-11"

    created = _post_attendance(
        client,
        class_id,
        date=day,
        student_name=student.name,
        type="病假",
        reason="发热就医",
        handled_note="已电话联系家长确认",
    ).json()["data"]

    result = _roll_call(
        client, class_id, day, [{"studentId": student.id, "type": "病假"}]
    ).json()["data"]
    assert result["changes"]["kept"] == 1
    assert result["changes"]["created"] == 0

    state = next(item for item in _day(client, class_id, day)["students"] if item["studentId"] == student.id)
    assert state["reason"] == "发热就医"
    assert state["handledNote"] == "已电话联系家长确认"

    rows = client.get("/api/v1/attendance", params={"classId": class_id, "q": "考勤留档甲"}).json()
    assert rows["meta"]["total"] == 1  # 没有多出一条
    assert rows["data"][0]["id"] == created["id"]  # 也没有换 id


def test_roll_call_resets_follow_up_when_the_type_changes(client, db_session):
    """类型改了，旧事由不再适用；跟进状态回到新类型的默认。"""
    class_id = _class_id(db_session)
    student = _student(db_session, "考勤改型甲", "A9009")
    day = "2026-11-12"

    _post_attendance(client, class_id, date=day, student_name=student.name, type="病假", reason="发热")
    result = _roll_call(
        client, class_id, day, [{"studentId": student.id, "type": "旷课"}]
    ).json()["data"]

    assert result["changes"]["updated"] == 1
    state = next(item for item in result["students"] if item["studentId"] == student.id)
    assert state["type"] == "旷课"
    assert state["handled"] == "待联系"  # 病假默认「无需联系」，改成旷课就该去联系家长
    assert state["reason"] == ""


def test_follow_up_default_follows_the_type(client, db_session):
    """跟进状态的默认值：旷课要联系家长，请假类家长已打过招呼。"""
    class_id = _class_id(db_session)
    cases = (
        ("考勤默认甲", "A9010", "旷课", "待联系"),
        ("考勤默认乙", "A9011", "病假", "无需联系"),
        ("考勤默认丙", "A9012", "事假", "无需联系"),
        ("考勤默认丁", "A9013", "迟到", "无需联系"),
    )
    for name, sno, kind, expected in cases:
        student = _student(db_session, name, sno)
        data = _post_attendance(
            client, class_id, date="2026-11-13", student_name=student.name, type=kind
        ).json()["data"]
        assert data["handled"] == expected, f"{kind} 的默认跟进状态应为 {expected}"


def test_confirmed_contact_survives_a_later_type_change(client, db_session):
    """「已通知」是既成事实，系统不悄悄撤掉。"""
    class_id = _class_id(db_session)
    student = _student(db_session, "考勤已通知甲", "A9014")
    day = "2026-11-14"

    created = _post_attendance(
        client, class_id, date=day, student_name=student.name, type="事假"
    ).json()["data"]
    client.patch(f"/api/v1/attendance/{created['id']}", json={"handled": "已通知"})

    result = _roll_call(
        client, class_id, day, [{"studentId": student.id, "type": "旷课"}]
    ).json()["data"]
    state = next(item for item in result["students"] if item["studentId"] == student.id)
    assert state["handled"] == "已通知"


# ---------- 点名：入参校验 ----------


def test_roll_call_rejects_unknown_type(client, db_session):
    class_id = _class_id(db_session)
    student = _student(db_session, "考勤类型甲", "A9015")

    response = _roll_call(
        client, class_id, "2026-11-15", [{"studentId": student.id, "type": "旷课了"}]
    )
    assert response.status_code == 400
    assert "旷课了" in response.json()["error"]["message"]


def test_roll_call_rejects_students_outside_the_class(client, db_session):
    """名单里有外人就报错，不静默忽略 —— 忽略等于老师以为记上了、其实没记。"""
    class_id = _class_id(db_session)
    student = _student(db_session, "考勤转出甲", "A9016")
    student.deleted_at = utcnow()
    db_session.commit()

    response = _roll_call(
        client, class_id, "2026-11-16", [{"studentId": student.id, "type": "病假"}]
    )
    assert response.status_code == 400
    assert "不在本班" in response.json()["error"]["message"]


def test_unknown_student_name_is_rejected(client, db_session):
    class_id = _class_id(db_session)
    response = _post_attendance(client, class_id, student_name="考勤查无此人")
    assert response.status_code == 400
    assert "考勤查无此人" in response.json()["error"]["message"]


def test_ambiguous_student_name_is_rejected(client, db_session):
    class_id = _class_id(db_session)
    _student(db_session, "考勤同名甲", "A9017")
    _student(db_session, "考勤同名甲", "A9018")

    response = _post_attendance(client, class_id, student_name="考勤同名甲")
    assert response.status_code == 400
    assert "都叫" in response.json()["error"]["message"]


def test_bad_date_is_reported_in_chinese(client, db_session):
    response = client.get(
        "/api/v1/attendance/day", params={"date": "上周五", "classId": _class_id(db_session)}
    )
    assert response.status_code == 400
    assert "日期" in response.json()["error"]["message"]


# ---------- 导入 / 导出 ----------


def test_export_carries_the_follow_up_state(client, db_session):
    class_id = _class_id(db_session)
    student = _student(db_session, "考勤导出甲", "A9019")
    _post_attendance(
        client,
        class_id,
        date="2026-11-17",
        student_name=student.name,
        type="旷课",
        reason="未到校",
        handled_note="给家长打电话未接",
    )

    response = client.get(
        "/api/v1/transfer/export/attendance.xlsx",
        params={"classId": class_id, "q": "考勤导出甲"},
    )
    sheet = load_workbook(io.BytesIO(response.content)).active
    headers = [cell.value for cell in sheet[1]]
    row = dict(zip(headers, next(iter(sheet.iter_rows(min_row=2, values_only=True)))))
    assert row["学生"] == "考勤导出甲"
    assert row["类型"] == "旷课"
    assert row["跟进状态"] == "待联系"
    assert row["处理情况"] == "给家长打电话未接"


def test_import_resolves_the_student_and_the_follow_up(client, db_session):
    class_id = _class_id(db_session)
    _student(db_session, "考勤导入甲", "A9020")
    csv_text = "日期,学生,类型,事由\n2026-11-18,考勤导入甲,事假,家中有事\n"

    preview = client.post(
        "/api/v1/transfer/import",
        params={"table": "attendance"},
        files={"file": ("a.csv", csv_text.encode("utf-8"), "text/csv")},
    ).json()["data"]
    assert preview["summary"]["problem"] == 0, preview

    values = [row["values"] for row in preview["rows"]]
    committed = client.post(
        "/api/v1/transfer/import/commit", params={"table": "attendance"}, json={"rows": values}
    ).json()["data"]
    assert committed["created"] == 1

    saved = client.get(
        "/api/v1/attendance", params={"classId": class_id, "q": "考勤导入甲"}
    ).json()["data"][0]
    assert saved["student_id"]  # 姓名解析成了 id，不是只留一个名字
    assert saved["handled"] == "无需联系"
